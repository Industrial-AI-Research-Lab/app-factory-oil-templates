import json

from .content import MISSING, display, headers, metadata, rows, text_report
from .csv_format import write_csv


def write_docx(path, data, style, family):
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = style["font"]
    normal.font.size = Pt(style["font_size"])
    doc.styles["Title"].font.color.rgb = RGBColor.from_string(
        style["accent"].lstrip("#")
    )
    doc.add_heading(style["heading"], 0)
    doc.add_heading(data["title"], 1)
    doc.add_paragraph(
        f"{data['project_id']} / {data['run_id']} / {data['result_id']} / {data['result_version']}"
    )
    for label, value in metadata(data, family):
        doc.add_paragraph(f"{label}: {value}")
    doc.add_paragraph(data["answer"])
    for table in data["tables"]:
        doc.add_heading(table["title"], 2)
        output = doc.add_table(rows=1, cols=len(table["columns"]) + 1)
        output.style = "Table Grid"
        for cell, value in zip(output.rows[0].cells, headers(table)):
            cell.text = value
        for values in rows(table):
            for cell, value in zip(output.add_row().cells, values):
                cell.text = display(value)
        if not table["rows"]:
            doc.add_paragraph(MISSING)
    doc.add_heading("Источники", 2)
    for source in data["sources"]:
        doc.add_paragraph(
            f"{source['id']} — {source['title']} — {display(source['url'])}"
        )
    if not data["sources"]:
        doc.add_paragraph(MISSING)
    doc.add_heading("Ограничения", 2)
    for value in data["limitations"] or [MISSING]:
        doc.add_paragraph(value)
    doc.save(path)
    Document(path)


def write_xlsx(path, data, style, family):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Результат"

    def append(target, values):
        if any(isinstance(value, str) and len(value) > 32767 for value in values):
            raise ValueError("XLSX text exceeds 32767 characters")
        target.append([MISSING if value is None else value for value in values])
        for cell in target[target.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    append(sheet, [style["heading"]])
    append(sheet, [data["title"]])
    append(
        sheet,
        [data["project_id"], data["run_id"], data["result_id"], data["result_version"]],
    )
    for label, value in metadata(data, family):
        append(sheet, [label, value])
    append(sheet, [data["answer"]])
    for table in data["tables"]:
        append(sheet, [table["title"]])
        append(sheet, headers(table))
        for values in rows(table):
            append(sheet, values)
        if not table["rows"]:
            append(sheet, [MISSING])
    append(sheet, ["Ограничения"])
    for limitation in data["limitations"] or [MISSING]:
        append(sheet, [limitation])
    sources = workbook.create_sheet("Источники")
    append(sources, ["ID", "Источник", "URL"])
    for source in data["sources"]:
        append(sources, [source["id"], source["title"], source["url"]])
    if not data["sources"]:
        append(sources, [MISSING])
    for target in workbook:
        target.freeze_panes = "A2"
        for column in range(1, target.max_column + 1):
            target.column_dimensions[get_column_letter(column)].width = style[
                "column_width"
            ]
        for cell in target[1]:
            cell.font = Font(bold=True, color=style["accent"].lstrip("#"))
    workbook.save(path)
    verified = load_workbook(path)
    try:
        for target in workbook:
            for row in target:
                for cell in row:
                    if cell.data_type == "n" and cell.value is not None:
                        actual = verified[target.title][cell.coordinate].value
                        if actual != cell.value:
                            raise ValueError(
                                f"XLSX cannot preserve numeric value at {target.title}!{cell.coordinate}"
                            )
    finally:
        verified.close()


def write_format(fmt, path, data, style, family, html, original):
    if fmt == "html":
        path.write_text(html(), encoding="utf-8")
    elif fmt == "pdf":
        from pypdf import PdfReader
        from weasyprint import HTML

        HTML(string=html()).write_pdf(path)
        if not PdfReader(path).pages:
            raise ValueError("PDF has no pages")
    elif fmt == "json":
        path.write_bytes(original)
        json.loads(path.read_text(encoding="utf-8"))
    elif fmt == "txt":
        path.write_text(text_report(data, style, family), encoding="utf-8")
    elif fmt == "csv":
        write_csv(path, data, style, family)
    elif fmt == "docx":
        write_docx(path, data, style, family)
    elif fmt == "xlsx":
        write_xlsx(path, data, style, family)
