import hashlib
import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


def test_small_typed_result_exports_real_seven_format_files(tmp_path):
    from reporting.render import render_report

    source = Path(__file__).parent / "fixtures" / "small.json"
    manifest = render_report(
        source,
        "knowledge",
        "1.0.0",
        ["txt", "json", "csv", "html", "docx", "xlsx", "pdf"],
        tmp_path,
    )
    folder = Path(manifest["manifest_path"]).parent
    assert manifest["result"] == {
        "id": "result-001",
        "version": "1",
        "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
    }
    assert json.loads((folder / "export-manifest.json").read_text()) == manifest
    assert (folder / "report-data.json").read_bytes() == source.read_bytes()
    for fmt, item in manifest["formats"].items():
        artifact = folder / item["filename"]
        assert item["status"] == "generated", (fmt, item)
        assert artifact.suffix == "." + fmt
        assert item["size_bytes"] == artifact.stat().st_size > 0
        assert item["sha256"] == hashlib.sha256(artifact.read_bytes()).hexdigest()
    assert json.loads((folder / "report.json").read_text()) == json.loads(
        source.read_text()
    )
    assert "&lt;без догадок&gt;" in (folder / "report.html").read_text()
    assert "Нет данных" in (folder / "report.txt").read_text()
    doc = Document(folder / "report.docx")
    assert any("Проверенный ответ" in p.text for p in doc.paragraphs)
    workbook = load_workbook(folder / "report.xlsx")
    assert workbook.sheetnames == ["Результат", "Источники"]
    cells = [cell for row in workbook["Результат"] for cell in row]
    assert any(cell.value == 12.5 and cell.data_type == "n" for cell in cells)
    assert any(
        cell.value == '=HYPERLINK("bad")' and cell.data_type == "s" for cell in cells
    )
    assert "Контрольный" in "".join(
        page.extract_text() for page in PdfReader(folder / "report.pdf").pages
    )
