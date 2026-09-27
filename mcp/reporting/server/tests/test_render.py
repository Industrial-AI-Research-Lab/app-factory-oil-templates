import hashlib
import io
import json
import time
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
import pytest


def fixture(name="small.json"):
    return json.loads((Path(__file__).parent / "fixtures" / name).read_text())


def test_small_result_renders_real_files():
    from app.rendering.render import render_bytes

    rendered = render_bytes(
        report_data=fixture(),
        template_id="knowledge",
        template_version="1.0.0",
        formats=["txt", "json", "csv", "html", "pdf", "docx", "xlsx"],
    )

    assert set(rendered.formats) == {
        "txt",
        "json",
        "csv",
        "html",
        "pdf",
        "docx",
        "xlsx",
    }
    for fmt, item in rendered.formats.items():
        assert item.status == "generated", (fmt, item.error)
        assert item.filename == f"report.{fmt}"
        assert item.payload
        assert item.size_bytes == len(item.payload)
        assert item.sha256 == hashlib.sha256(item.payload).hexdigest()
        assert item.error is None

    assert json.loads(rendered.formats["json"].payload) == fixture()
    assert (
        b"&lt;\xd0\xb1\xd0\xb5\xd0\xb7 \xd0\xb4\xd0\xbe\xd0\xb3\xd0\xb0\xd0\xb4\xd0\xbe\xd0\xba&gt;"
        in rendered.formats["html"].payload
    )
    assert "Нет данных" in rendered.formats["txt"].payload.decode()

    doc = Document(io.BytesIO(rendered.formats["docx"].payload))
    assert any("Проверенный ответ" in paragraph.text for paragraph in doc.paragraphs)
    workbook = load_workbook(io.BytesIO(rendered.formats["xlsx"].payload))
    assert workbook.sheetnames == ["Результат", "Источники"]
    cells = [cell for row in workbook["Результат"] for cell in row]
    assert any(cell.value == 12.5 and cell.data_type == "n" for cell in cells)
    assert any(
        cell.value == '=HYPERLINK("bad")' and cell.data_type == "s" for cell in cells
    )
    pages = PdfReader(io.BytesIO(rendered.formats["pdf"].payload)).pages
    assert "Контрольный" in "".join(page.extract_text() for page in pages)


def test_template_version_changes_appearance_without_changing_result():
    from app.rendering.render import render_bytes

    data = fixture()
    first = render_bytes(data, "knowledge", "1.0.0", ["html", "json"])
    second = render_bytes(data, "knowledge", "1.1.0", ["html", "json"])

    assert first.result_sha256 == second.result_sha256
    assert first.formats["json"].payload == second.formats["json"].payload
    assert first.formats["html"].payload != second.formats["html"].payload
    assert first.template_snapshot != second.template_snapshot


@pytest.mark.parametrize("template_id", ["news", "knowledge"])
@pytest.mark.parametrize("template_version", ["1.0.0", "1.1.0"])
def test_published_versions_keep_legacy_html_template_snapshot(
    template_id, template_version
):
    from app.rendering.render import render_bytes

    rendered = render_bytes(
        fixture("news.json" if template_id == "news" else "small.json"),
        template_id,
        template_version,
        ["html"],
    )

    assert rendered.template_snapshot["template.html.j2"] == (
        "10856662d56f2d65e455fa11645fb1ac10d98bd33d284e6ad256ada49a654567"
    )


def test_repeated_ooxml_renders_have_stable_bytes():
    from app.rendering.render import render_bytes

    data = fixture()
    first = render_bytes(data, "knowledge", "1.0.0", ["docx", "xlsx"])
    time.sleep(2.1)
    second = render_bytes(data, "knowledge", "1.0.0", ["docx", "xlsx"])

    for fmt in ("docx", "xlsx"):
        assert first.formats[fmt].payload == second.formats[fmt].payload
        assert first.formats[fmt].sha256 == second.formats[fmt].sha256


@pytest.mark.parametrize(
    "template_id,template_version,fixture_name",
    [
        ("knowledge", "1.0.0", "small.json"),
        ("knowledge", "1.1.0", "small.json"),
        ("knowledge", "1.2.0", "small.json"),
        ("news", "1.0.0", "news.json"),
        ("news", "1.1.0", "news.json"),
    ],
)
def test_each_template_version_renders(template_id, template_version, fixture_name):
    from app.rendering.render import render_bytes

    rendered = render_bytes(
        fixture(fixture_name), template_id, template_version, ["html"]
    )

    assert rendered.template == {"id": template_id, "version": template_version}
    assert rendered.formats["html"].status == "generated"


@pytest.mark.parametrize("template_id", ["news", "knowledge"])
def test_empty_result_renders_all_formats(template_id):
    from app.rendering.render import render_bytes

    data = fixture("empty.json")
    if template_id == "news":
        data["metadata"] = {
            "topic": "Новости",
            "period": "2026-09-01",
            "conditions": "Проверенные",
            "publication_count": 0,
            "event_count": 0,
            "categories": [],
        }
    rendered = render_bytes(
        data,
        template_id,
        "1.0.0",
        ["txt", "json", "csv", "html", "pdf", "docx", "xlsx"],
    )

    assert all(item.status == "generated" for item in rendered.formats.values())
    assert "Нет данных" in rendered.formats["html"].payload.decode()


def test_long_table_keeps_last_row_in_pdf_and_docx():
    from app.rendering.render import render_bytes

    data = fixture("long.json")
    rendered = render_bytes(data, "knowledge", "1.0.0", ["pdf", "docx", "xlsx"])

    assert all(item.status == "generated" for item in rendered.formats.values())
    pages = PdfReader(io.BytesIO(rendered.formats["pdf"].payload)).pages
    assert len(pages) > 1
    assert "Строка 119" in "".join(page.extract_text() for page in pages)
    doc = Document(io.BytesIO(rendered.formats["docx"].payload))
    assert "Строка 119" in doc.tables[0].rows[-1].cells[0].text


def test_one_format_failure_does_not_remove_successful_formats():
    from app.rendering.render import render_bytes

    data = fixture()
    data["answer"] = "x" * 32768
    rendered = render_bytes(data, "knowledge", "1.0.0", ["txt", "xlsx"])

    assert rendered.formats["txt"].status == "generated"
    assert rendered.formats["txt"].payload
    assert rendered.formats["xlsx"].status == "failed"
    assert rendered.formats["xlsx"].payload is None
    assert rendered.formats["xlsx"].error == {
        "code": "RENDER_FAILED",
        "stage": "render",
        "message": "Не удалось сформировать файл",
    }


@pytest.mark.parametrize(
    ("cap_delta", "expected_status"),
    [(-1, "failed"), (0, "generated"), (1, "generated")],
)
def test_format_size_limit_has_an_inclusive_boundary(
    monkeypatch, cap_delta, expected_status
):
    from app.rendering import render as rendering

    data = fixture("small.json")
    json_size = len(
        json.dumps(
            data,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            allow_nan=False,
        ).encode("utf-8")
    )
    monkeypatch.setattr(
        rendering,
        "MAX_FORMAT_BYTES",
        json_size + cap_delta,
        raising=False,
    )
    monkeypatch.setattr(
        rendering,
        "MAX_TOTAL_OUTPUT_BYTES",
        json_size * 2,
        raising=False,
    )

    result = rendering.render_bytes(data, "knowledge", "1.0.0", ["json"])

    assert result.formats["json"].status == expected_status
    if expected_status == "failed":
        assert result.formats["json"].error["code"] == "OUTPUT_TOO_LARGE"


def test_aggregate_size_limit_rejects_only_the_format_that_crosses_it(monkeypatch):
    from app.rendering import render as rendering

    data = fixture("small.json")
    json_size = len(
        json.dumps(
            data,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            allow_nan=False,
        ).encode("utf-8")
    )
    monkeypatch.setattr(rendering, "MAX_FORMAT_BYTES", json_size * 2, raising=False)
    monkeypatch.setattr(rendering, "MAX_TOTAL_OUTPUT_BYTES", json_size, raising=False)

    result = rendering.render_bytes(
        data,
        "knowledge",
        "1.0.0",
        ["json", "txt"],
    )

    assert result.formats["json"].status == "generated"
    assert result.formats["txt"].status == "failed"
    assert result.formats["txt"].error["code"] == "OUTPUT_TOO_LARGE"
